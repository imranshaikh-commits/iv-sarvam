"""The real renderer stack, not mocks (audit 2026-10-12, finding 12).

The rest of the suite mocks LibreOffice and skips D2 when the binary is
absent, so a green run on a laptop or a bare CI runner said nothing about
whether the production image can draw a diagram or produce a PDF. These run
the actual binaries. Outside the brain image they skip; inside it
(SHILPI_REQUIRE_STACK=1, set by CI's image job and by deploy/deploy.sh) a
missing binary is a FAILURE, because that image is what ships.
"""
import io
import os
import shutil

import pytest
from docx import Document
from PIL import Image

import diagram_engine
import export_engine
import rfp_intake

REQUIRED = os.environ.get("SHILPI_REQUIRE_STACK") == "1"


def _need(binary: str):
    if shutil.which(binary):
        return
    if REQUIRED:
        pytest.fail(f"{binary} is missing from the production image")
    pytest.skip(f"{binary} not installed here")


def _spec():
    return diagram_engine.DiagramSpec.model_validate({
        "diagram_type": "architecture", "title": "Smoke",
        "nodes": [{"id": "a", "label": "Identity Provider", "group": "IAM"},
                  {"id": "b", "label": "Business Application", "group": "Apps"}],
        "edges": [{"source": "a", "target": "b", "label": "SAML"}]})


def test_d2_draws_a_real_diagram():
    _need("d2")
    _need("rsvg-convert")
    png = diagram_engine.render_spec(_spec(), fmt="png", renderer="d2")
    assert png and Image.open(io.BytesIO(png)).width >= 600


def test_graphviz_fallback_draws_a_real_diagram():
    _need("dot")
    png = diagram_engine.render_spec(_spec(), fmt="png", renderer="graphviz")
    assert png and Image.open(io.BytesIO(png)).width > 50


def _docx() -> bytes:
    doc = Document()
    doc.add_heading("Smoke", level=1)
    doc.add_paragraph("A proposal paragraph.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_libreoffice_turns_a_docx_into_a_pdf_and_poppler_reads_it(tmp_path):
    _need("soffice")
    pdf, err = export_engine.export_pdf(_docx())
    assert err is None and pdf and pdf.startswith(b"%PDF")
    _need("pdftoppm")
    path = tmp_path / "smoke.pdf"
    path.write_bytes(pdf)
    pages = rfp_intake.rasterise(str(path), str(tmp_path / "pages"))
    assert len(pages) == 1


def _mixed_docx() -> bytes:
    """Page 1: a page of text. Page 2: nothing but a picture (a scanned annexure)."""
    from docx.shared import Inches
    doc = Document()
    doc.add_paragraph("Section 4. Access management requirements for the tender. " * 12)
    doc.add_page_break()
    pic = io.BytesIO()
    Image.new("RGB", (600, 800), "white").save(pic, "PNG")
    pic.seek(0)
    doc.add_picture(pic, width=Inches(5))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_poppler_counts_pages_as_they_are_and_a_scanned_page_is_read_as_an_image(tmp_path):
    """Deep audit F12, with the real binaries. pdftotext ends every page with a
    form feed, so a plain split counted one page too many; and a page with no
    text layer inside a text PDF has to reach the vision reader, with the right
    page's image."""
    import asyncio
    _need("soffice")
    _need("pdftotext")
    _need("pdftoppm")
    pdf, err = export_engine.export_pdf(_mixed_docx())
    assert err is None and pdf
    path = tmp_path / "mixed.pdf"
    path.write_bytes(pdf)
    text_pages = rfp_intake.read_text_layer(str(path))
    assert len(text_pages) == 2                                   # not 3
    assert not rfp_intake.page_is_thin(text_pages[0]) and rfp_intake.page_is_thin(text_pages[1])
    assert not rfp_intake.needs_vision(text_pages)                # the average hides page 2

    seen = []

    class _Empty:
        field_values, requirements, eligibility_gates, structure_headings = [], [], [], []

    async def fake_structured(model, messages):
        content = messages[1]["content"]
        seen.append("vision" if isinstance(content, list) else "text")
        return _Empty()

    ex = asyncio.run(rfp_intake.extract_rfp(
        str(path), "mixed.pdf", [], fake_structured, str(tmp_path / "pages")))
    assert seen == ["text", "vision"]
    assert ex.pages_total == 2 and ex.vision_pages == [2] and ex.unread_pages == []
