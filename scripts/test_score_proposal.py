"""Offline checks for score_proposal.py on a synthetic document."""
import io
import itertools
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import score_proposal as S  # noqa: E402
from docx import Document  # noqa: E402
from PIL import Image  # noqa: E402


def _png(w=40, h=20):
    b = io.BytesIO()
    Image.new("RGB", (w, h), "white").save(b, "PNG")
    b.seek(0)
    return b


def _doc(path, exec_words=450, repeats=3, gantt=True, extra=""):
    d = Document()
    d.add_heading("Executive Summary", level=1)
    # Varied 5-word sentences: repeated filler is (correctly) flagged degenerate.
    nouns = ["identity", "access", "policy", "audit", "role", "request", "review",
             "account", "session", "vault", "token", "report", "workflow"]
    verbs = ["covers", "governs", "routes", "records", "approves", "checks", "links"]
    sents = [f"The {a} {v} each {b}." for a, v, b in
             itertools.islice(itertools.product(nouns, verbs, nouns), exec_words // 5)]
    d.add_paragraph(" ".join(sents))
    d.add_heading("Scope", level=1)
    d.add_heading("Detail", level=2)
    d.add_paragraph(" ".join(["There are 5,000 users."] * repeats) + " [SME REVIEW] " + extra)
    t = d.add_table(rows=2, cols=4)
    for i, h in enumerate(["Workstream", "Phase", "Start Week", "End Week"]):
        t.rows[0].cells[i].text = h
    for i, v in enumerate(["IGA", "Build", "1", "8"]):
        t.rows[1].cells[i].text = v
    if gantt:
        from docx.shared import Inches
        d.add_picture(_png(900, 450), width=Inches(3))   # 300 dpi: sharp
    d.add_paragraph("Price: To be confirmed")
    d.save(path)


def test_measure_counts_what_a_reviewer_counts():
    with tempfile.TemporaryDirectory() as tmp:
        p = f"{tmp}/a.docx"
        _doc(p)
        m = S.measure(p)
    assert m["exec_summary_words"] == 450
    assert m["tables"] == 1 and m["images"] == 1 and m["gantt_charts"] == 1
    assert m["sme"] == 1 and m["tbc"] == 1
    assert m["top_repeated_figures"][0] == ("5,000", 3)
    assert m["headings"] == {"H1": 2, "H2": 1}


def test_gates_catch_the_esnad_0925_failures():
    with tempfile.TemporaryDirectory() as tmp:
        p = f"{tmp}/a.docx"
        _doc(p, exec_words=185, repeats=57, gantt=False, extra="Azure AD Amlak")
        m = S.measure(p)
    cfg = {"must_mention": ["Power BI"], "forbid": ["Azure AD"], "other_clients": ["Amlak"]}
    failed = [g for g, ok, _ in S.gates(m, None, cfg) if not ok]
    for gate in ("exec_summary_words", "max_figure_repeats", "no other client",
                 "named systems", "no forbidden", "gantt"):
        assert any(f.startswith(gate) for f in failed), (gate, failed)


def test_clean_document_passes_its_hygiene_gates_and_exit_code_follows():
    with tempfile.TemporaryDirectory() as tmp:
        p = f"{tmp}/a.docx"
        _doc(p)
        m = S.measure(p)
        hygiene = [g for g in S.gates(m, None, {}) if not g[0].startswith(
            ("prose_words", "table_words", "tables", "images"))]
        assert all(ok for _, ok, _ in hygiene), hygiene
        assert S.main([p]) == 1          # absolute size floors fail a 1-page doc


def test_blurry_image_fails_the_gate():
    """Imran, 2026-09-27: no blurry image anywhere. 40 px printed 2 in wide."""
    from docx.shared import Inches
    with tempfile.TemporaryDirectory() as tmp:
        p = f"{tmp}/a.docx"
        d = Document()
        d.add_picture(_png(), width=Inches(2))
        d.save(p)
        m = S.measure(p)
    assert m["blurry_images"] == 1
    assert any(g == "no blurry images" and not ok for g, ok, _ in S.gates(m, None, {}))


def test_term_matching_ignores_spacing():
    assert S._mentions("PowerBI dashboards", "Power BI")
    assert not S._mentions("Power users", "Power BI")


def test_size_gates_scale_to_a_small_baseline():
    # A scoped deal's own IV document must pass the size gates it sets.
    with tempfile.TemporaryDirectory() as tmp:
        p = f"{tmp}/a.docx"
        _doc(p, exec_words=100)
        m = S.measure(p)
    size = [g for g in S.gates(m, m, {}) if g[0].startswith(
        ("prose_words", "table_words", "tables", "images", "exec_summary"))]
    assert size and all(ok for _, ok, _ in size), size
